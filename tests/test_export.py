"""
=========================================================
Music Collection Manager
Export Tests
=========================================================

Milestone 5C (1/N)

pytest tests for core/export.py.

Every test writes to a pytest tmp_path - none of these touch the
real project's exports/ folder. Content is read back and checked,
not just file existence, since a file that "exists" could still
have the wrong data.
"""

from __future__ import annotations

import csv
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from core.export import (
    _sanitised_sheet_title,
    export_to_csv,
    export_to_excel,
    export_to_pdf,
)

SAMPLE_ROWS = [
    {"SongID": 1, "Title": "Song A", "BPM": 128.0, "Year": None},
    {"SongID": 2, "Title": "Song B", "BPM": None, "Year": 1999},
]
SAMPLE_COLUMNS = ["SongID", "Title", "BPM", "Year"]


# ============================================================
# CSV
# ============================================================


def test_csv_header_matches_columns(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    export_to_csv(SAMPLE_ROWS, SAMPLE_COLUMNS, path)

    with open(path, newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader)

    assert header == SAMPLE_COLUMNS


def test_csv_row_content_matches_source_data(tmp_path: Path) -> None:
    path = tmp_path / "export.csv"
    export_to_csv(SAMPLE_ROWS, SAMPLE_COLUMNS, path)

    with open(path, newline="", encoding="utf-8") as handle:
        rows = list(csv.DictReader(handle))

    assert len(rows) == 2
    assert rows[0]["Title"] == "Song A"
    assert rows[0]["BPM"] == "128.0"
    assert rows[1]["Year"] == "1999"


def test_csv_ignores_row_keys_not_in_columns(tmp_path: Path) -> None:
    rows = [{"SongID": 1, "Title": "Song A", "Extra": "should not appear"}]
    path = tmp_path / "export.csv"
    export_to_csv(rows, ["SongID", "Title"], path)

    with open(path, newline="", encoding="utf-8") as handle:
        reader = csv.reader(handle)
        header = next(reader)

    assert header == ["SongID", "Title"]


def test_csv_handles_empty_rows(tmp_path: Path) -> None:
    path = tmp_path / "empty.csv"
    export_to_csv([], SAMPLE_COLUMNS, path)

    with open(path, newline="", encoding="utf-8") as handle:
        rows = list(csv.reader(handle))

    assert rows == [SAMPLE_COLUMNS]


# ============================================================
# Excel
# ============================================================


def test_excel_header_matches_columns_and_is_bold(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    export_to_excel(SAMPLE_ROWS, SAMPLE_COLUMNS, path)

    workbook = load_workbook(path)
    worksheet = workbook.active
    assert worksheet is not None

    header = [cell.value for cell in worksheet[1]]
    assert header == SAMPLE_COLUMNS
    assert all(cell.font.bold for cell in worksheet[1])


def test_excel_row_content_matches_source_data(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    export_to_excel(SAMPLE_ROWS, SAMPLE_COLUMNS, path)

    workbook = load_workbook(path)
    worksheet = workbook.active
    assert worksheet is not None

    first_row = [cell.value for cell in worksheet[2]]
    assert first_row == [1, "Song A", 128.0, None]


def test_excel_sheet_title_is_applied(tmp_path: Path) -> None:
    path = tmp_path / "export.xlsx"
    export_to_excel(
        SAMPLE_ROWS, SAMPLE_COLUMNS, path, sheet_title="My Songs Report"
    )

    workbook = load_workbook(path)
    assert workbook.active is not None
    assert workbook.active.title == "My Songs Report"


def test_excel_handles_empty_rows(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    export_to_excel([], SAMPLE_COLUMNS, path)

    workbook = load_workbook(path)
    worksheet = workbook.active
    assert worksheet is not None
    assert worksheet.max_row == 1
    assert [cell.value for cell in worksheet[1]] == SAMPLE_COLUMNS


# ============================================================
# PDF
# ============================================================


def test_pdf_contains_title_and_data(tmp_path: Path) -> None:
    path = tmp_path / "export.pdf"
    export_to_pdf(SAMPLE_ROWS, SAMPLE_COLUMNS, path, title="Song Report")

    text = PdfReader(path).pages[0].extract_text()

    assert "Song Report" in text
    assert "Song A" in text
    assert "Song B" in text


def test_pdf_handles_empty_rows(tmp_path: Path) -> None:
    path = tmp_path / "empty.pdf"
    export_to_pdf([], SAMPLE_COLUMNS, path, title="Empty Report")

    text = PdfReader(path).pages[0].extract_text()

    assert "Empty Report" in text
    for column in SAMPLE_COLUMNS:
        assert column in text


def test_pdf_wide_table_stays_on_one_page(tmp_path: Path) -> None:
    # 12 columns matches Records - the widest real table in this app.
    # A naive layout could overflow the page width; colWidths are
    # divided evenly to guarantee it never does.
    wide_columns = [f"Col{i}" for i in range(12)]
    wide_rows = [
        {column: f"val{i}" for column in wide_columns} for i in range(3)
    ]
    path = tmp_path / "wide.pdf"

    export_to_pdf(wide_rows, wide_columns, path, title="Wide Table")

    reader = PdfReader(path)
    assert len(reader.pages) == 1
    text = reader.pages[0].extract_text()
    assert "Col0" in text
    assert "Col11" in text


# ============================================================
# _sanitised_sheet_title
# ============================================================


def test_sanitised_sheet_title_removes_invalid_characters() -> None:
    assert _sanitised_sheet_title("Songs/Artists?") == "Songs_Artists_"


def test_sanitised_sheet_title_truncates_to_31_characters() -> None:
    long_title = "A" * 50
    result = _sanitised_sheet_title(long_title)
    assert len(result) == 31


def test_sanitised_sheet_title_falls_back_when_empty() -> None:
    assert _sanitised_sheet_title("") == "Export"
    assert _sanitised_sheet_title("   ") == "Export"


def test_sanitised_sheet_title_leaves_valid_titles_unchanged() -> None:
    assert _sanitised_sheet_title("Songs Without Artists") == "Songs Without Artists"
