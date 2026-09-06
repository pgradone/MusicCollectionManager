"""
=========================================================
Music Collection Manager
Export
=========================================================

Milestone 5C (1/N)

CSV/Excel/PDF export - "Add global search, advanced filtering,
statistics, and exports" from the original specification's Search &
Reports phase.

Deliberately stateless and independent of DatabaseContext, like
core/backup.py and core/settings.py: every function here takes
already-fetched tabular data (a list of row dicts plus the column
order to use) and writes it to a file. main.py decides what data to
export (the Browse grid, a Report, or Search results) - this module
only knows how to write it out in three formats.
"""

from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Any


def export_to_csv(
    rows: list[dict[str, Any]],
    columns: list[str],
    path: Path,
) -> None:
    """Write rows to a CSV file, using columns as the header and column order."""

    with open(path, "w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for row in rows:
            writer.writerow({column: row.get(column) for column in columns})


def export_to_excel(
    rows: list[dict[str, Any]],
    columns: list[str],
    path: Path,
    *,
    sheet_title: str = "Export",
) -> None:
    """
    Write rows to an .xlsx workbook, one sheet, with a bold header
    row and column widths sized to fit their content.
    """

    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    worksheet = workbook.active
    assert worksheet is not None
    worksheet.title = _sanitised_sheet_title(sheet_title)

    worksheet.append(columns)
    for cell in worksheet[1]:
        cell.font = Font(bold=True)

    for row in rows:
        worksheet.append([row.get(column) for column in columns])

    for col_index, column in enumerate(columns, start=1):
        content_lengths = [len(str(column))]
        content_lengths.extend(
            len(str(row.get(column)))
            for row in rows
            if row.get(column) is not None
        )
        width = min(max(content_lengths) + 2, 60)
        worksheet.column_dimensions[get_column_letter(col_index)].width = width

    workbook.save(path)


def export_to_pdf(
    rows: list[dict[str, Any]],
    columns: list[str],
    path: Path,
    *,
    title: str = "Export",
) -> None:
    """
    Write rows to a landscape-oriented PDF table, with a title and a
    repeating header row on every page. Column widths are divided
    evenly across the usable page width, so even a wide table (many
    columns) always fits rather than overflowing the page.
    """

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import landscape, letter
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    page_width, page_height = landscape(letter)
    left_margin = right_margin = 36
    usable_width = page_width - left_margin - right_margin
    column_width = usable_width / max(len(columns), 1)

    document = SimpleDocTemplate(
        str(path),
        pagesize=landscape(letter),
        leftMargin=left_margin,
        rightMargin=right_margin,
    )
    styles = getSampleStyleSheet()

    data = [columns] + [
        ["" if row.get(column) is None else str(row.get(column)) for column in columns]
        for row in rows
    ]

    table = Table(data, colWidths=[column_width] * len(columns), repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#dddddd")),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ]
        )
    )

    story = [Paragraph(title, styles["Title"]), Spacer(1, 12), table]
    document.build(story)


def _sanitised_sheet_title(title: str) -> str:
    """
    Excel worksheet names can't contain \\ / ? * [ ] and are capped
    at 31 characters.
    """

    cleaned = re.sub(r"[\\/?*\[\]:]", "_", title).strip()

    return cleaned[:31] if cleaned else "Export"
